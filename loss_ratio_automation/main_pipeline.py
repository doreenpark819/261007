"""
손해율 월말 집계 및 워드 리포트 자동화 — 메인 파이프라인

  1) 로데이터 + 담보 마스터 Left Join (담보코드 기준, VLOOKUP 대체)
  2) UY 누적 / CY 당월 손해율 산출 (당해 + 전년 동기)
  3) 담보별 전년 동기 대비 위험보험료 증감률(%) · 손해율 증감(%p)
  4) Top 20 추출 (위험보험료 상위 / 손해율 상위)
  5) MS Word(.docx) 리포트 출력

실행:  python main_pipeline.py [--base 25.09] [--uy-start 25.06]
Phase 2(내부망) 전환 시에는 아래 CONFIG 블록(또는 환경변수)만 수정한다.
"""
from __future__ import annotations

import argparse
import os
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# =============================================================================
# CONFIG — Phase 2에서는 이 블록만 바꾸면 된다 (환경변수로도 덮어쓰기 가능)
# =============================================================================
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("LR_DATA_DIR", BASE_DIR / "data"))
OUTPUT_DIR = Path(os.environ.get("LR_OUTPUT_DIR", BASE_DIR / "output"))
RAW_FILE = DATA_DIR / os.environ.get("LR_RAW_FILE", "raw_data.csv")
MASTER_FILE = DATA_DIR / os.environ.get("LR_MASTER_FILE", "code_master.csv")
REPORT_FILE = OUTPUT_DIR / "monthly_loss_report.docx"

BASE_YM = "25.09"        # 마감 기준월 (CY 당월 = 이 달, UY 누적 = 이 달까지)
UY_START_YM = "25.06"    # UY 누적 대상 계약년월 시작
TOP_N = 20
# 손해율 Top N 은 위험보험료가 작은 담보의 극단값을 피하기 위해
# 전체 위험보험료 대비 비중이 이 값 이상인 담보만 대상으로 한다 (0 이면 제한 없음)
LR_RANK_MIN_PREMIUM_SHARE = 0.0005   # 0.05%
AMOUNT_UNIT = 1_000_000              # 금액 표시 단위 (원) → 백만원
AMOUNT_UNIT_LABEL = "백만원"

# 원본 컬럼명 매핑 (내부망 실제 컬럼명이 다르면 여기만 수정)
COL = {
    "contract_ym": "계약년월",
    "close_ym": "마감년월",
    "main_prod": "주상품코드",
    "rider_prod": "특약상품코드",
    "elapsed": "경과개월수",
    "uw_type": "심사유형(일반.간편)",
    "code": "담보코드",
    "premium": "위험보험료 합계",
    "paid": "지급보험금합계",
    "reserve_chg": "지급비금변동분 합계",
}
M_NAME, M_BENEFIT = "담보명", "대급부명"
UNMAPPED = "(미매핑)"

# 손해액 = 지급보험금 + 지급비금변동분,  손해율 = 손해액 / 위험보험료 × 100
LOSS_COMPONENTS = ("paid", "reserve_chg")

# 워드 스타일
FONT_NAME = "맑은 고딕"
HEADER_FILL = "1F3864"     # 표 헤더 배경 (남색)
HEADER_FONT = RGBColor(0xFF, 0xFF, 0xFF)
ZEBRA_FILL = "F2F5FA"
BORDER_COLOR = "A6B3C9"
UP_COLOR = RGBColor(0xC0, 0x00, 0x00)    # 악화(손해율↑) 빨강
DOWN_COLOR = RGBColor(0x1F, 0x4E, 0xA8)  # 개선(손해율↓) 파랑


# =============================================================================
# 1. 데이터 로드 & 조인
# =============================================================================
def read_csv(path: Path, **kw) -> pd.DataFrame:
    """내부망 CSV는 CP949인 경우가 많으므로 UTF-8 → CP949 순으로 시도."""
    for enc in ("utf-8-sig", "cp949"):
        try:
            return pd.read_csv(path, encoding=enc, **kw)
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError(f"인코딩을 판별할 수 없습니다: {path}")


def normalize_ym(s: pd.Series) -> pd.Series:
    """'25.06', '2025.06', '202506', '2025-06' 등을 pandas Period(M)로 통일."""
    digits = s.astype(str).str.replace(r"\D", "", regex=True)
    digits = digits.where(digits.str.len() != 4, "20" + digits)   # YYMM → 20YYMM
    return pd.PeriodIndex(digits.str[:6], freq="M")


def to_period(ym: str) -> pd.Period:
    return normalize_ym(pd.Series([ym]))[0]


def fmt_ym(p: pd.Period) -> str:
    return f"{p.year % 100:02d}.{p.month:02d}"


def load_data() -> pd.DataFrame:
    str_cols = {COL[k]: str for k in ("contract_ym", "close_ym", "main_prod", "rider_prod", "code")}
    raw = read_csv(RAW_FILE, dtype=str_cols)
    master = read_csv(MASTER_FILE, dtype={COL["code"]: str})

    missing = [c for c in COL.values() if c not in raw.columns]
    if missing:
        raise KeyError(f"로데이터에 필요한 컬럼이 없습니다: {missing}")

    raw[COL["code"]] = raw[COL["code"]].str.strip()
    master[COL["code"]] = master[COL["code"]].str.strip()
    dup = master[COL["code"]].duplicated()
    if dup.any():
        print(f"[경고] 마스터 담보코드 중복 {dup.sum()}건 → 첫 행만 사용")
        master = master[~dup]

    for k in ("premium", "paid", "reserve_chg"):
        raw[COL[k]] = pd.to_numeric(
            raw[COL[k]].astype(str).str.replace(",", ""), errors="coerce").fillna(0)

    df = raw.merge(master[[COL["code"], M_NAME, M_BENEFIT]], on=COL["code"], how="left")
    unmapped = df[M_NAME].isna()
    if unmapped.any():
        codes = sorted(df.loc[unmapped, COL["code"]].unique())
        print(f"[경고] 마스터 미매핑 담보코드 {len(codes)}개 ({unmapped.sum()}행): {codes[:10]}")
        df[[M_NAME, M_BENEFIT]] = df[[M_NAME, M_BENEFIT]].fillna(UNMAPPED)

    df["_cym"] = normalize_ym(df[COL["contract_ym"]])
    df["_clym"] = normalize_ym(df[COL["close_ym"]])
    df["_prem"] = df[COL["premium"]]
    df["_loss"] = sum(df[COL[k]] for k in LOSS_COMPONENTS)
    print(f"[OK] 로데이터 {len(raw):,}행 / 마스터 {len(master):,}건 조인 완료")
    return df


# =============================================================================
# 2. UY / CY 필터 & 집계
# =============================================================================
@dataclass
class Window:
    label: str
    cur: pd.DataFrame
    prev: pd.DataFrame
    cur_desc: str
    prev_desc: str


def uy_filter(df: pd.DataFrame, start: pd.Period, base: pd.Period) -> pd.DataFrame:
    """UY 누적: 계약년월 ∈ [start, base] 인 계약의 마감 base 까지 누적 실적."""
    return df[(df["_cym"] >= start) & (df["_cym"] <= base) & (df["_clym"] <= base)]


def cy_filter(df: pd.DataFrame, base: pd.Period) -> pd.DataFrame:
    """CY 당월: 마감년월 == base 인 한 달 실적 (계약년월 무관)."""
    return df[df["_clym"] == base]


def build_windows(df: pd.DataFrame, base: pd.Period, uy_start: pd.Period) -> list[Window]:
    pb, ps = base - 12, uy_start - 12
    windows = [
        Window("UY 누적", uy_filter(df, uy_start, base), uy_filter(df, ps, pb),
               f"{fmt_ym(uy_start)}~{fmt_ym(base)} 계약, {fmt_ym(base)} 마감 누적",
               f"{fmt_ym(ps)}~{fmt_ym(pb)} 계약, {fmt_ym(pb)} 마감 누적"),
        Window("CY 당월", cy_filter(df, base), cy_filter(df, pb),
               f"{fmt_ym(base)} 마감월", f"{fmt_ym(pb)} 마감월"),
    ]
    for w in windows:
        if w.cur.empty:
            raise ValueError(f"{w.label} 당해 데이터가 없습니다 ({w.cur_desc})")
        if w.prev.empty:
            print(f"[경고] {w.label} 전년 동기 데이터가 없습니다 ({w.prev_desc}) → 증감 공란")
    return windows


def loss_ratio(loss, prem):
    return (loss / prem * 100).where(prem != 0) if isinstance(prem, pd.Series) \
        else (loss / prem * 100 if prem else float("nan"))


def totals(d: pd.DataFrame) -> dict:
    p, l = d["_prem"].sum(), d["_loss"].sum()
    return {"prem": p, "loss": l, "lr": loss_ratio(l, p)}


def compare(cur: pd.DataFrame, prev: pd.DataFrame, keys: list[str]) -> pd.DataFrame:
    """keys 단위로 당해/전년 집계 후 증감 계산."""
    agg = lambda d: d.groupby(keys, dropna=False)[["_prem", "_loss"]].sum()
    out = agg(cur).join(agg(prev), how="outer", rsuffix="_py").fillna(0).reset_index()
    out["lr"] = loss_ratio(out["_loss"], out["_prem"])
    out["lr_py"] = loss_ratio(out["_loss_py"], out["_prem_py"])
    out["prem_chg_pct"] = ((out["_prem"] / out["_prem_py"] - 1) * 100).where(out["_prem_py"] != 0)
    out["lr_chg_p"] = out["lr"] - out["lr_py"]
    return out


def by_code(w: Window) -> pd.DataFrame:
    return compare(w.cur, w.prev, [COL["code"], M_BENEFIT, M_NAME])


def top_n(tbl: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    cur = tbl[tbl["_prem"] > 0]
    by_prem = cur.nlargest(TOP_N, "_prem")
    eligible = cur[cur["_prem"] >= cur["_prem"].sum() * LR_RANK_MIN_PREMIUM_SHARE]
    by_lr = eligible.nlargest(TOP_N, "lr")
    return by_prem, by_lr


# =============================================================================
# 3. Word 리포트
# =============================================================================
def set_cell_fill(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_table_borders(table, color: str = BORDER_COLOR, size: int = 4) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size if edge.startswith("inside") else size * 2))
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_font(run, size=9, bold=False, color: RGBColor | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def write_cell(cell, text: str, align="left", bold=False, color=None, size=9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_font(p.add_run(text), size=size, bold=bold, color=color)


def amt(v) -> str:
    return "-" if pd.isna(v) else f"{v / AMOUNT_UNIT:,.0f}"


def pct(v) -> str:
    return "-" if pd.isna(v) else f"{v:,.1f}%"


def pct_signed(v) -> str:
    return "-" if pd.isna(v) else f"{v:+,.1f}%"


def pt_signed(v) -> str:
    return "-" if pd.isna(v) else f"{v:+,.1f}%p"


def chg_color(v):
    """손해율 증감: 악화(+) 빨강, 개선(-) 파랑."""
    if pd.isna(v) or abs(v) < 0.05:
        return None
    return UP_COLOR if v > 0 else DOWN_COLOR


def add_table(doc, headers, rows, aligns, widths_cm, colors=None):
    """rows: list[list[str]], colors: list[list[RGBColor|None]] (optional)."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        set_cell_fill(hdr.cells[i], HEADER_FILL)
        write_cell(hdr.cells[i], h, "center", bold=True, color=HEADER_FONT)
    for r_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = colors[r_idx][i] if colors else None
            write_cell(cells[i], val, aligns[i], color=c)
            if r_idx % 2 == 1:
                set_cell_fill(cells[i], ZEBRA_FILL)
    # 열 너비: Word/LibreOffice 모두 반영되도록 tblGrid + 셀 너비 + 고정 레이아웃 지정
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for i, col in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        col.set(qn("w:w"), str(Cm(widths_cm[i]).twips))
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    return t


def heading(doc, text, level=1, new_page=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), size=14 if level == 1 else 11.5, bold=True,
             color=RGBColor(0x1F, 0x38, 0x64))
    return p


def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text), size=8.5, color=RGBColor(0x59, 0x59, 0x59))


def summary_section(doc, windows: list[Window]):
    heading(doc, "1. 핵심 요약 (전사)")
    headers = ["구분", "기간", f"위험보험료({AMOUNT_UNIT_LABEL})", f"손해액({AMOUNT_UNIT_LABEL})",
               "손해율", "전년 손해율", "위험보험료 증감률", "손해율 증감"]
    rows, colors = [], []
    for w in windows:
        c, p = totals(w.cur), totals(w.prev)
        prem_chg = (c["prem"] / p["prem"] - 1) * 100 if p["prem"] else float("nan")
        lr_chg = c["lr"] - p["lr"]
        rows.append([w.label, w.cur_desc, amt(c["prem"]), amt(c["loss"]), pct(c["lr"]),
                     pct(p["lr"]), pct_signed(prem_chg), pt_signed(lr_chg)])
        colors.append([None] * 7 + [chg_color(lr_chg)])
    add_table(doc, headers, rows,
              ["center", "left", "right", "right", "right", "right", "right", "right"],
              [2.0, 5.6, 3.3, 3.0, 2.0, 2.2, 3.4, 2.6], colors)
    note(doc, f"※ 전년 동기: UY {windows[0].prev_desc} / CY {windows[1].prev_desc}")


def breakdown_section(doc, w: Window, key: str, title: str):
    heading(doc, title, level=2)
    tbl = compare(w.cur, w.prev, [key]).sort_values("_prem", ascending=False)
    headers = [key, f"위험보험료({AMOUNT_UNIT_LABEL})", f"손해액({AMOUNT_UNIT_LABEL})",
               "손해율", "위험보험료 증감률", "손해율 증감(p)"]
    rows = [[str(r[key]), amt(r["_prem"]), amt(r["_loss"]), pct(r["lr"]),
             pct_signed(r["prem_chg_pct"]), pt_signed(r["lr_chg_p"])] for _, r in tbl.iterrows()]
    colors = [[None] * 5 + [chg_color(r["lr_chg_p"])] for _, r in tbl.iterrows()]
    add_table(doc, headers, rows, ["left"] + ["right"] * 5, [4.0, 3.6, 3.6, 2.6, 3.4, 3.2], colors)


def top_table(doc, tbl: pd.DataFrame, title: str, new_page=False):
    heading(doc, title, level=2, new_page=new_page)
    headers = ["순위", "담보코드", "대급부명", "담보명", f"위험보험료({AMOUNT_UNIT_LABEL})",
               "위험보험료 증감률", "손해율(%)", "전동기 대비 증감(p)"]
    rows, colors = [], []
    for rank, (_, r) in enumerate(tbl.iterrows(), start=1):
        rows.append([str(rank), r[COL["code"]], r[M_BENEFIT], r[M_NAME], amt(r["_prem"]),
                     pct_signed(r["prem_chg_pct"]), pct(r["lr"]), pt_signed(r["lr_chg_p"])])
        colors.append([None] * 7 + [chg_color(r["lr_chg_p"])])
    add_table(doc, headers, rows,
              ["center", "left", "left", "left", "right", "right", "right", "right"],
              [1.1, 2.0, 2.2, 7.4, 3.2, 3.0, 2.4, 3.2], colors)


def build_report(windows: list[Window], base: pd.Period, out_path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(9)

    # 타이틀 & 메타데이터
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"{base.year}.{base.month:02d} 기준 손해율 결산 리포트"), size=20, bold=True,
             color=RGBColor(0x1F, 0x38, 0x64))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"마감 기준월 {base.year}.{base.month:02d}  |  작성일 {date.today():%Y.%m.%d}  |  "
                       f"금액 단위 {AMOUNT_UNIT_LABEL}"), size=9, color=RGBColor(0x59, 0x59, 0x59))
    note(doc, "※ 손해율 = (지급보험금 + 지급비금변동분) ÷ 위험보험료 × 100,  "
              "증감(p) = 당해 손해율 − 전년 동기 손해율 (빨강: 악화 / 파랑: 개선)")

    summary_section(doc, windows)

    for i, w in enumerate(windows, start=2):
        heading(doc, f"{i}. {w.label} 손해율 ({w.cur_desc})", new_page=i > 2)
        breakdown_section(doc, w, COL["uw_type"], f"{i}-1. 심사유형별")
        breakdown_section(doc, w, M_BENEFIT, f"{i}-2. 대급부별")
        by_prem, by_lr = top_n(by_code(w))
        top_table(doc, by_prem, f"{i}-3. 위험보험료 상위 {TOP_N} 담보", new_page=True)
        top_table(doc, by_lr, f"{i}-4. 손해율 상위 {TOP_N} 담보", new_page=True)
        if LR_RANK_MIN_PREMIUM_SHARE:
            note(doc, f"※ 손해율 상위는 위험보험료 비중 {LR_RANK_MIN_PREMIUM_SHARE:.2%} 이상 담보 대상")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


# =============================================================================
# main
# =============================================================================
def main() -> None:
    ap = argparse.ArgumentParser(description="손해율 월말 집계 및 워드 리포트 생성")
    ap.add_argument("--base", default=BASE_YM, help="마감 기준월 (예: 25.09)")
    ap.add_argument("--uy-start", default=UY_START_YM, help="UY 누적 계약년월 시작 (예: 25.06)")
    ap.add_argument("--out", default=str(REPORT_FILE), help="출력 docx 경로")
    args = ap.parse_args()

    base, uy_start = to_period(args.base), to_period(args.uy_start)
    if uy_start > base:
        raise ValueError("--uy-start 는 --base 보다 늦을 수 없습니다")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    df = load_data()
    windows = build_windows(df, base, uy_start)

    for w in windows:
        c, p = totals(w.cur), totals(w.prev)
        print(f"[{w.label}] 위험보험료 {c['prem']:,.0f} / 손해액 {c['loss']:,.0f} / "
              f"손해율 {c['lr']:.2f}% (전년 {p['lr']:.2f}%, {c['lr'] - p['lr']:+.2f}%p)")
        # 검증용 담보별 집계 CSV
        by_code(w).rename(columns={"_prem": "위험보험료", "_loss": "손해액",
                                   "_prem_py": "전년_위험보험료", "_loss_py": "전년_손해액",
                                   "lr": "손해율", "lr_py": "전년_손해율",
                                   "prem_chg_pct": "위험보험료_증감률", "lr_chg_p": "손해율_증감p"}) \
            .to_csv(OUTPUT_DIR / f"by_code_{w.label.split()[0]}_{fmt_ym(base)}.csv",
                    index=False, encoding="utf-8-sig")

    build_report(windows, base, Path(args.out))
    print(f"[OK] 리포트 생성: {args.out}")


if __name__ == "__main__":
    main()

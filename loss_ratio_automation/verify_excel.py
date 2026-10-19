"""
[외부망 전용] 엑셀 보고서 검증 — 재계산된 xlsx의 '보고서' 시트 숫자를 pandas 로 독립 재계산해 대조한다.

  python verify_excel.py excel/손해율_보고서.xlsx      (먼저 LibreOffice/엑셀로 재계산된 파일이어야 함)
"""
import sys

import numpy as np
import openpyxl
import pandas as pd

TOL = 1e-9


def prev_ym(x):
    return x - 89 if x % 100 == 1 else x - 1


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    raw = pd.DataFrame(list(wb["raw_data"].values))
    raw.columns = raw.iloc[0]
    raw = raw.iloc[1:].reset_index(drop=True)
    mst = pd.DataFrame(list(wb["code_master"].values)).iloc[:, :5]
    mst.columns = ["담보코드", "담보명", "대급부명", "구분", "고당대통"]
    mst = mst.iloc[1:].dropna(subset=["담보코드"]).fillna("")
    run = wb["실행"]
    cfg = {k: run[c].value for k, c in [("base", "C4"), ("launch", "C5"), ("top_lr", "C6"), ("over", "C7"),
                                         ("top_prem", "C8"), ("min_share", "C9"), ("godang", "C10")]}
    cfg["order1"] = [run.cell(11, j).value for j in range(3, 13) if run.cell(11, j).value]
    cfg["order3"] = [run.cell(12, j).value for j in range(3, 13) if run.cell(12, j).value]
    rpt = [list(r) for r in wb["보고서"].iter_rows(values_only=True)]
    return raw, mst, cfg, rpt


def main(path):
    raw, mst, cfg, rpt = load(path)
    base, launch = int(cfg["base"]), int(cfg["launch"])
    prev = prev_ym(base)
    uys = base // 100 * 100 + 1
    cs, ce = prev_ym(uys), prev
    a = raw["계약년월"].astype(int)
    b = raw["마감년월"].astype(int)
    raw = raw.assign(P=raw["위험보험료 합계"].astype(float),
                     L=raw["지급보험금합계"].astype(float) + raw["지급비금변동분 합계"].astype(float),
                     C=pd.to_numeric(raw["건수"], errors="coerce").fillna(0))
    W = {
        "CY": b == base, "CYP": b == prev, "ALL": (b >= launch) & (b <= base), "ALLP": (b >= launch) & (b <= prev),
        "UM": (a >= uys) & (a <= base) & (b == base), "UMC": (a >= cs) & (a <= ce) & (b == ce),
        "UC": (a >= uys) & (a <= base) & (b <= base), "UCC": (a >= cs) & (a <= ce) & (b <= ce),
    }
    tot = {k: (raw.loc[m, "P"].sum(), raw.loc[m, "L"].sum()) for k, m in W.items()}
    lr = lambda k: tot[k][1] / tot[k][0]

    errs = []

    def chk(label, got, exp):
        if exp is None or (isinstance(exp, float) and np.isnan(exp)):
            ok = got in (None, "")
        elif isinstance(exp, str):
            ok = got == exp
        else:
            ok = got is not None and got != "" and abs(float(got) - exp) <= TOL * max(1, abs(exp))
        if not ok:
            errs.append(f"{label}: 엑셀={got!r} 기대={exp!r}")

    tbl_rows = [i for i, r in enumerate(rpt) if isinstance(r[0], str) and r[0].startswith("TBL")]
    # 표: CY
    r = tbl_rows[0]
    for col, k in ((3, "CY"), (4, "CYP")):
        chk(f"CY {k} 손해율", rpt[r + 1][col], lr(k))
        chk(f"CY {k} 위험P", rpt[r + 2][col], tot[k][0])
        chk(f"CY {k} 손해액", rpt[r + 3][col], tot[k][1])
    chk("CY 증감 손해율", rpt[r + 1][5], lr("CY") - lr("CYP"))
    chk("CY 누적 위험P", rpt[r + 2][6], tot["ALL"][0])
    chk("CY 누적 손해액", rpt[r + 3][6], tot["ALL"][1])
    chk("CY 누적 텍스트", rpt[r + 1][6], f"{lr('ALL'):.0%}\n(*~{prev // 100 % 100}.{prev % 100}월: {lr('ALLP'):.0%})")
    # 표: UY
    r = tbl_rows[1]
    for col, k in ((3, "UM"), (4, "UMC"), (6, "UC"), (7, "UCC")):
        chk(f"UY {k} 손해율", rpt[r + 2][col], lr(k))
        chk(f"UY {k} 위험P", rpt[r + 3][col], tot[k][0])
        chk(f"UY {k} 손해액", rpt[r + 4][col], tot[k][1])
    chk("UY 당월 比", rpt[r + 2][5], lr("UM") - lr("UMC"))
    chk("UY 누적 比", rpt[r + 2][8], lr("UC") - lr("UCC"))

    # 급부 단위
    m = mst.copy()
    m["grp"] = np.where(m["구분"].str.strip() == "", m["대급부명"], m["구분"])
    m["gd"] = (m["고당대통"].str.upper().str.strip() == "Y").astype(int)
    d = raw.merge(m, on="담보코드", how="inner")
    d = d[d["담보명"].str.strip() != ""]
    first = m[m["담보명"] != ""].drop_duplicates("담보명").set_index("담보명")
    g = pd.DataFrame(index=first.index)
    da = d["계약년월"].astype(int)
    db = d["마감년월"].astype(int)
    DW = {"UC": (da >= uys) & (da <= base) & (db <= base), "UCC": (da >= cs) & (da <= ce) & (db <= ce),
          "ALL": (db >= launch) & (db <= base), "ALLP": (db >= launch) & (db <= prev)}
    for k, mk in DW.items():
        s = d[mk].groupby("담보명")[["P", "L"]].sum()
        g[k + "P"] = s["P"].reindex(g.index).fillna(0)
        g[k + "L"] = s["L"].reindex(g.index).fillna(0)
    g["CNT"] = d[DW["ALL"]].groupby("담보명")["C"].sum().reindex(g.index).fillna(0)
    g["grp"] = first["grp"]
    g["big"] = first["대급부명"]
    g["gd"] = m.groupby("담보명")["gd"].max().reindex(g.index)
    div = lambda x, y: np.where(y != 0, x / np.where(y != 0, y, 1), np.nan)
    g["lr"], g["lrc"] = div(g.UCL, g.UCP), div(g.UCCL, g.UCCP)
    g["diff"] = g.lr - g.lrc
    g["lra"], g["lrap"] = div(g.ALLL, g.ALLP), div(g.ALLPL, g.ALLPP)
    g["pf"] = g.ALLP / g.ALLP.sum()
    g["elig"] = (g.UCP > 0) & (g.UCP >= float(cfg["min_share"] or 0) * g.UCP.sum())
    g["rk_all"] = g.lr.where(g.elig).rank(ascending=False, method="first")
    g["rk_ex"] = g.lr.where(g.elig & (g.gd == 0)).rank(ascending=False, method="first")
    o1 = {v: i + 1 for i, v in enumerate(cfg["order1"])}
    o3 = {v: i + 1 for i, v in enumerate(cfg["order3"])}
    in1 = (g.rk_ex <= int(cfg["top_lr"])) | (g.gd == 1)
    g["g1"] = np.where(g.gd == 1, cfg["godang"], g.grp)
    t1 = g[in1].assign(k=lambda x: x.g1.map(o1).fillna(50) * 100000 + x.rk_all.fillna(99999)).sort_values("k")
    t2 = g[~in1 & (g.lr > float(cfg["over"]))].assign(
        k=lambda x: x.grp.map(o1).fillna(50) * 100000 + x.rk_all.fillna(99999)).sort_values("k")
    g["rk_p"] = g.UCP.where(g.UCP > 0).rank(ascending=False, method="first")
    t3 = g[g.rk_p <= int(cfg["top_prem"])].assign(k=lambda x: x.big.map(o3).fillna(50) * 100000 + x.rk_p).sort_values("k")

    def body(ti, head, n_expected):
        r0 = tbl_rows[ti] + head
        rows = []
        for rr in rpt[r0:]:
            if rr[2] in (None, ""):
                break
            rows.append(rr)
        if len(rows) != n_expected:
            errs.append(f"표{ti - 1} 행 수: 엑셀={len(rows)} 기대={n_expected}")
        return rows

    for ti, t, grpcol, cols in ((2, t1, "g1", ("lr", "lrc", "diff", "lra")), (3, t2, "grp", ("lr", "lrc", "diff", "lra")),
                                (4, t3, "big", ("ALLP", "pf", "CNT", "lra"))):
        head = 1 if ti == 4 else 2
        rows = body(ti, head, len(t))
        for rr, (nm, e) in zip(rows, t.iterrows()):
            chk(f"표{ti - 1} 급부명", rr[2], nm)
            chk(f"표{ti - 1} {nm} 구분", rr[1], e[grpcol])
            for j, c in enumerate(cols):
                chk(f"표{ti - 1} {nm} {c}", rr[3 + j], None if pd.isna(e[c]) else float(e[c]))
        if ti == 2:
            def c1(e):
                new = not pd.isna(e["diff"]) and e["diff"] > 0
                cum = not pd.isna(e.lra) and not pd.isna(e.lrap) and e.lra > e.lrap
                return 1 if new and not cum else 2 if cum and not new else 0
            col_exp = {nm: c1(e) for nm, e in t.iterrows()}
        elif ti == 3:
            col_exp = {nm: (1 if e["diff"] > 0 else 2 if e["diff"] < 0 else 0) for nm, e in t.iterrows()}
        else:
            col_exp = {}
        for rr in rows:
            if col_exp and rr[2] in col_exp:
                chk(f"표{ti - 1} {rr[2]} 색코드", rr[9], col_exp[rr[2]])
        if ti == 2:  # 합계 행
            f = tbl_rows[ti] + head + 40
            sel = t
            chk("주요급부 평균 UY", rpt[f][3], sel.UCL.sum() / sel.UCP.sum())
            chk("주요급부 평균 비교", rpt[f][4], sel.UCCL.sum() / sel.UCCP.sum())
            chk("주요급부 평균 누적", rpt[f][6], sel.ALLL.sum() / sel.ALLP.sum())
            chk("전급부 평균 UY", rpt[f + 1][3], lr("UC"))
            chk("전급부 평균 누적", rpt[f + 1][6], lr("ALL"))
            chk("비중 UY", rpt[f + 2][3], sel.UCP.sum() / tot["UC"][0])
            chk("비중 누적", rpt[f + 2][6], sel.ALLP.sum() / tot["ALL"][0])

    print(f"기준월 {base} / 비교 계약 {cs}~{ce} / 표(1) {len(t1)}행, 표(2) {len(t2)}행, 표(3) {len(t3)}행")
    print(f"CY {lr('CY'):.1%} (전월 {lr('CYP'):.1%}) / UY당월 {lr('UM'):.1%} vs {lr('UMC'):.1%} / "
          f"UY누적 {lr('UC'):.1%} vs {lr('UCC'):.1%} / 누적 {lr('ALL'):.1%}")
    if errs:
        print(f"[불일치 {len(errs)}건]")
        print("\n".join(errs[:40]))
        sys.exit(1)
    print("[OK] 보고서 시트 전 항목 일치")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "excel/손해율_보고서.xlsx")

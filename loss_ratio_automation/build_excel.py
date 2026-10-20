"""
[외부망 전용] 엑셀 보고서 템플릿 생성기 — 내부망에는 결과 .xlsx / .bas 만 반입한다.

  excel/손해율_보고서.xlsx : 수식 기반 집계 + 워드 보고서 레이아웃 (더미 데이터 포함)
  excel/report_macro.bas   : 데이터 불러오기 + 워드 생성 VBA (CP949)

시트
  실행        : 기준월 입력(1칸) · 보고서 설정 · 자동 계산 구간 · 데이터 점검
  보고서      : 워드와 같은 순서·모양의 보고서 (A열 = 워드 출력 태그)
  급부별집계  : 급부(담보명) 단위 집계 · 순위 · 표 선정
  담보별집계  : 담보코드 단위 집계 (로데이터 SUMIFS)
  raw_data    : 로데이터 붙여넣기 (A~K, K=건수 선택)
  code_master : 담보코드 | 담보명(=급부) | 대급부명(표3 구분) | 구분(표1·2 구분, 빈칸=대급부명) | 고당대통(Y)
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

BASE = Path(__file__).resolve().parent
OUT = BASE / "excel" / "손해율_보고서.xlsx"

MAX_CODES = 1000                 # 담보 마스터 수용 한도
MAX_GB = 500                     # 급부 종류 수용 한도
SLOTS = {"T1": 40, "T2": 60, "T3": 30}

FONT = "맑은 고딕"
F = lambda size=10, bold=False, color="000000": Font(name=FONT, size=size, bold=bold, color=color)
FILL_HDR = PatternFill("solid", fgColor="D9D9D9")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_SET = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

NF_LR = '0%;"△"0%;0%'                       # 손해율
NF_PT = '0%"p";"△"0%"p";0%"p"'              # 손해율 증감
NF_AMT = '#,##0,,;"△"#,##0,,;0'             # 원 → 백만원
NF_SH = '0.0%;"△"0.0%;0.0%'                 # 비중
NF_SH_PT = '0.0%"p";"△"0.0%"p";0.0%"p"'
NF_CNT = '#,##0;"△"#,##0;0'

RAW, MST, CD, GB, RPT = "raw_data", "code_master", "담보별집계", "급부별집계", "보고서"
R = f"{RAW}!"
C_ = f"'{CD}'!"
G_ = f"'{GB}'!"
CL = MAX_CODES + 1
GL = MAX_GB + 1

# ---------------------------------------------------------------------------
# 집계 구간 (SUMIFS 조건) — A=계약년월, B=마감년월 (202608 같은 숫자)
# ---------------------------------------------------------------------------
CRIT = {
    "CY":   f'{R}$B:$B,BASE_YM',                                              # CY 당월
    "CYP":  f'{R}$B:$B,PREV_YM',                                              # CY 전월
    "ALL":  f'{R}$B:$B,">="&LAUNCH_YM,{R}$B:$B,"<="&BASE_YM',                 # 판매개시~기준월 누적
    "ALLP": f'{R}$B:$B,">="&LAUNCH_YM,{R}$B:$B,"<="&PREV_YM',                 # 판매개시~전월 누적
    "UM":   f'{R}$A:$A,">="&UY_START,{R}$A:$A,"<="&BASE_YM,{R}$B:$B,BASE_YM',          # UY 당월
    "UMC":  f'{R}$A:$A,">="&CMP_START,{R}$A:$A,"<="&CMP_END,{R}$B:$B,CMP_END',         # UY 당월 비교
    "UC":   f'{R}$A:$A,">="&UY_START,{R}$A:$A,"<="&BASE_YM,{R}$B:$B,"<="&BASE_YM',     # UY 누적
    "UCC":  f'{R}$A:$A,">="&CMP_START,{R}$A:$A,"<="&CMP_END,{R}$B:$B,"<="&CMP_END',    # UY 누적 비교
}


def prem(w, extra=""):
    return f"SUMIFS({R}$H:$H,{CRIT[w]}{extra})"


def loss(w, extra=""):
    return f"(SUMIFS({R}$I:$I,{CRIT[w]}{extra})+SUMIFS({R}$J:$J,{CRIT[w]}{extra}))"


def lr(w):
    return f'IFERROR({loss(w)}/{prem(w)},"")'


yy = lambda x: f"MOD(INT({x}/100),100)"
# COUNTIF/SUMIF/MATCH 조건용: 이름에 ~ * ? 가 있어도 문자 그대로 비교되도록 이스케이프
esc = lambda x: f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"~","~~"),"*","~*"),"?","~?")'
mm = lambda x: f"MOD({x},100)"


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def widths(ws, ws_widths, start=1):
    for i, w in enumerate(ws_widths, start=start):
        ws.column_dimensions[get_column_letter(i)].width = w


def cell(ws, ref, value, font=None, fill=None, align=None, fmt=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font or F()
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    if border:
        c.border = border
    return c


def box(ws, rng, fill=None, font=None, align=CENTER):
    """범위 전체에 테두리/정렬(필요시 채움·글꼴) 적용."""
    for row in ws[rng]:
        for c in row:
            c.border = BORDER
            c.alignment = align
            if fill:
                c.fill = fill
            if font:
                c.font = font


def merge(ws, rng, value=None):
    ws.merge_cells(rng)
    if value is not None:
        ws[rng.split(":")[0]] = value


# ---------------------------------------------------------------------------
# 실행 시트
# ---------------------------------------------------------------------------
def sheet_run(wb):
    ws = wb.active
    ws.title = "실행"
    widths(ws, [3, 34, 14, 12, 12, 12, 12, 12, 12, 12, 12, 12])
    ws.sheet_view.showGridLines = False
    cell(ws, "B1", "손해율 월말 보고서", F(16, True, "1F3864"))
    cell(ws, "B3", "① 입력 (매월 노란 칸 하나만 바꾸면 됨)", F(11, True, "1F3864"))

    inputs = [  # 행, 항목, 값, 서식, 이름, 설명
        (4, "마감 기준월", 202608, "0", None, "예) 202608"),
        (5, "최초판매월 (누적 시작)", 202506, "0", None, "보통 바꿀 일 없음"),
        (6, "표(1) 손해율 상위 N", 20, "0", "TOP_LR_N", "주요급부 = UY 손해율 상위 N + 고당대통"),
        (7, "표(2) 손해율 초과 기준", 1, "0%", "OVER_LR", "그외 이 값 초과 급부"),
        (8, "표(3) 위험보험료 상위 N", 15, "0", "TOP_PREM_N", "최대 30"),
        (9, "손해율 순위 대상 최소 보험료 비중", 0, "0.00%", "MIN_SHARE", "보험료가 아주 작은 급부 제외용 (0이면 제한 없음)"),
        (10, "고당대통 표시 구분명", "기타", "@", "GODANG_GROUP", "표(1)에서 고당대통 급부를 묶는 구분명"),
    ]
    for r, label, val, fmt, nm, desc in inputs:
        cell(ws, f"B{r}", label, F(10, True))
        cell(ws, f"C{r}", val, F(12 if r == 4 else 10, r == 4, "0000FF" if r == 4 else "000000"),
             FILL_INPUT if r == 4 else FILL_SET, CENTER, fmt, BORDER)
        cell(ws, f"E{r}", desc, F(9, color="595959"))
        if nm:
            name(wb, nm, f"실행!$C${r}")
    for r, label, items, nm in ((11, "표(1)(2) 구분 표시 순서", ["사망", "진단", "장해", "수술", "입원", "기타"], "ORDER1"),
                                (12, "표(3) 급부 표시 순서", ["암이외진단", "암진단", "장해", "수술", "입원", "기타", "사망"], "ORDER3")):
        cell(ws, f"B{r}", label, F(10, True))
        for j in range(10):
            col = get_column_letter(3 + j)
            cell(ws, f"{col}{r}", items[j] if j < len(items) else None, F(9), FILL_SET, CENTER, None, BORDER)
        name(wb, nm, f"실행!$C${r}:$L${r}")
    cell(ws, "B13", "  ※ 순서에 없는 구분은 맨 뒤에 표시. 구분 안에서는 손해율(표3은 보험료) 큰 순", F(9, color="595959"))

    # 내부 숫자값 (텍스트로 입력돼도 숫자로)
    grey = F(8, color="A6A6A6")
    for r, nm in ((4, "BASE_YM"), (5, "LAUNCH_YM")):
        cell(ws, f"D{r}", f"=IFERROR(--TRIM(C{r}),0)", grey, fmt="0")
        name(wb, nm, f"실행!$D${r}")

    cell(ws, "B15", "② 자동 계산 구간 (확인용)", F(11, True, "1F3864"))
    prev = lambda x: f"IF(MOD({x},100)=1,{x}-89,{x}-1)"
    derived = [  # 행, 항목, 수식, 서식, 이름
        (16, "전월 (CY 비교 · UY 비교 마감월)", f"={prev('BASE_YM')}", "0", "PREV_YM"),
        (17, "UY 계약 시작월 (기준연도 1월)", "=INT(BASE_YM/100)*100+1", "0", "UY_START"),
        (18, "UY 비교 계약 시작월 (1개월 전)", f"={prev('UY_START')}", "0", "CMP_START"),
        (19, "UY 비교 계약 종료월 (= 전월)", "=PREV_YM", "0", "CMP_END"),
        (20, "UY 개월수 (직전 N개월)", "=MOD(BASE_YM,100)", "0", "N_MON"),
        (21, "라벨: 기준월", f'="\'"&{yy("BASE_YM")}&"."&{mm("BASE_YM")}', "@", "L_BASE"),
        (22, "라벨: 전월", f'="\'"&{yy("PREV_YM")}&"."&{mm("PREV_YM")}', "@", "L_PREV"),
        (23, "라벨: UY 당기", f'="\'"&{yy("UY_START")}&"."&{mm("UY_START")}&"~"&{mm("BASE_YM")}', "@", "L_UYC"),
        (24, "라벨: UY 비교", f'=IF(CMP_START=CMP_END,"\'"&{yy("CMP_START")}&"."&{mm("CMP_START")},'
                             f'"\'"&{yy("CMP_START")}&"."&{mm("CMP_START")}&"~\'"&{yy("CMP_END")}&"."&{mm("CMP_END")})', "@", "L_CMP"),
        (25, "라벨: 최초판매월", f'="\'"&{yy("LAUNCH_YM")}&"."&{mm("LAUNCH_YM")}', "@", "L_LAUNCH"),
        (26, "라벨: UY 연도", f"={yy('BASE_YM')}", "0", "YYN"),
    ]
    for r, label, f, fmt, nm in derived:
        cell(ws, f"B{r}", label, F(9))
        cell(ws, f"C{r}", f, F(10, True), None, CENTER, fmt, BORDER)
        name(wb, nm, f"실행!$C${r}")

    cell(ws, "B28", "③ 매월 작업 순서", F(11, True, "1F3864"))
    steps = [
        "1) 노란 칸에 마감 기준월(예: 202608) 입력 → 구간·라벨·표 제목은 모두 자동",
        "2) 매크로 [LoadRawData] 실행 → 로데이터 파일 선택 (기존 데이터 자동 삭제·형식 정리)",
        "     매크로를 못 쓰면: raw_data 시트 내용 삭제 → A1에 헤더 포함 붙여넣기 (A~J 원본 10개 열 + K열 건수)",
        "     ※ 로데이터는 최초판매월 마감분부터 기준월 마감분까지 전체 필요",
        "3) 아래 ④ 데이터 점검의 종합 상태 확인 → '보고서' 시트에서 숫자 확인",
        "4) 매크로 [MakeWordReport] 실행 → 이 파일과 같은 폴더에 손해율_보고서_YYYYMM.docx 생성",
        "담보 마스터가 바뀐 경우에만 code_master 시트 교체 (최대 1,000개). 급부 = 담보명 기준으로 합산",
    ]
    for i, t in enumerate(steps, start=29):
        cell(ws, f"B{i}", t, F(10))

    # ④ 데이터 점검
    top = 37
    cell(ws, f"B{top}", "④ 데이터 점검 (자동)", F(11, True, "1F3864"))
    for j, h in enumerate(["점검 항목", "결과", "상태", "문제일 때 조치"]):
        cell(ws, f"{'BCDE'[j]}{top + 1}", h, F(9, True), FILL_HDR, CENTER, None, BORDER)
    ws.merge_cells(f"E{top + 1}:L{top + 1}")
    rows_ = f"(COUNTA({R}$G:$G)-1)"
    ok = lambda cond, bad="오류": f'=IF({cond},"정상","{bad}")'
    checks = [
        ("기준월 입력 형식", "=BASE_YM", ok("AND(BASE_YM>=190001,BASE_YM<=209912,MOD(BASE_YM,100)>=1,MOD(BASE_YM,100)<=12)"),
         "202608처럼 6자리 숫자로 입력"),
        ("최초판매월 형식", "=LAUNCH_YM", ok("AND(LAUNCH_YM>=190001,MOD(LAUNCH_YM,100)>=1,MOD(LAUNCH_YM,100)<=12,LAUNCH_YM<=BASE_YM)"),
         "6자리 숫자, 기준월 이전"),
        ("로데이터 행 수", f"={rows_}", ok(f"{rows_}>0"), "raw_data에 데이터 없음 → LoadRawData 실행"),
        ("로데이터 열 수 (10~11)", f"=COUNTA({R}$1:$1)", ok(f"AND(COUNTA({R}$1:$1)>=10,COUNTA({R}$1:$1)<=11)", "확인"),
         "A~J 원본 10개 열 + K열 건수(선택) 순서 확인"),
        ("건수(K열) 입력 여부", f"=COUNT({R}$K:$K)", ok(f"COUNT({R}$K:$K)>0", "확인"), "없으면 표(3) 건수가 0으로 표시"),
        ("년월이 숫자가 아닌 행", f"=2*{rows_}-COUNT({R}$A:$A)-COUNT({R}$B:$B)",
         ok(f"2*{rows_}-COUNT({R}$A:$A)-COUNT({R}$B:$B)=0"), "매크로가 자동 변환 (수동: 데이터→텍스트 나누기→마침)"),
        ("금액이 숫자가 아닌 칸", f"=3*{rows_}-COUNT({R}$H:$H)-COUNT({R}$I:$I)-COUNT({R}$J:$J)",
         ok(f"3*{rows_}-COUNT({R}$H:$H)-COUNT({R}$I:$I)-COUNT({R}$J:$J)=0"), "위와 같음 (빈칸은 0)"),
        ("년월 범위 (열 순서 확인)", f'=MIN({R}$A:$B)&" ~ "&MAX({R}$A:$B)',
         ok(f"AND(MIN({R}$A:$B)>=190001,MAX({R}$A:$B)<=209912)"), "A열=계약년월, B열=마감년월인지 확인"),
        ("기준월 마감 행 수", f"=COUNTIF({R}$B:$B,BASE_YM)", ok(f"COUNTIF({R}$B:$B,BASE_YM)>0"), "기준월 또는 데이터 기간 확인"),
        ("전월 마감 행 수", f"=COUNTIF({R}$B:$B,PREV_YM)", ok(f"COUNTIF({R}$B:$B,PREV_YM)>0", "확인"),
         "없으면 전월·동기간 비교가 공란"),
        ("최초판매월 이전 마감 행 (누적 제외)", f'=COUNTIF({R}$B:$B,"<"&LAUNCH_YM)', ok(f'COUNTIF({R}$B:$B,"<"&LAUNCH_YM)=0', "확인"),
         "최초판매월 입력값 확인"),
        ("기준월 이후 마감 행 (집계 제외)", f'=COUNTIF({R}$B:$B,">"&BASE_YM)', ok(f'COUNTIF({R}$B:$B,">"&BASE_YM)=0', "확인"),
         "과거 월 재작성이면 정상"),
        ("마스터에 없는 담보코드 행", f"={rows_}-SUM({C_}$P$2:$P${CL})", ok(f"{rows_}-SUM({C_}$P$2:$P${CL})=0", "확인"),
         "전사 합계엔 포함, 급부별 표에선 빠짐 → code_master에 추가"),
        ("마스터 담보코드 중복", f"=SUM({C_}$Q$2:$Q${CL})", ok(f"SUM({C_}$Q$2:$Q${CL})=0"), "code_master 중복 제거"),
        ("마스터 담보명(급부) 빈칸", f'=COUNTIFS({MST}!$A$2:$A${CL},"?*",{MST}!$B$2:$B${CL},"")',
         ok(f'COUNTIFS({MST}!$A$2:$A${CL},"?*",{MST}!$B$2:$B${CL},"")=0', "확인"), "급부명 없는 담보는 급부별 표에서 빠짐"),
        ("마스터 담보 수 (최대 1,000)", f"=COUNTA({MST}!$A:$A)-1", ok(f"COUNTA({MST}!$A:$A)-1<={MAX_CODES}"), "초과 시 템플릿 확장 필요"),
        ("급부 종류 수 (최대 500)", f"=MAX({C_}$O$2:$O${CL})", ok(f"MAX({C_}$O$2:$O${CL})<={MAX_GB}"), "초과 시 템플릿 확장 필요"),
        ("표(1) 행 수 (최대 40)", f"=COUNT({G_}$AB$2:$AB${GL})", ok(f"COUNT({G_}$AB$2:$AB${GL})<={SLOTS['T1']}"), "상위 N 또는 고당대통 수 확인"),
        ("표(2) 행 수 (최대 60)", f"=COUNT({G_}$AF$2:$AF${GL})", ok(f"COUNT({G_}$AF$2:$AF${GL})<={SLOTS['T2']}"), "초과 기준 손해율 확인"),
        ("표(3) 행 수 (최대 30)", f"=COUNT({G_}$AK$2:$AK${GL})", ok(f"COUNT({G_}$AK$2:$AK${GL})<={SLOTS['T3']}"), "상위 N 확인"),
    ]
    r = top + 2
    first = r
    for label, val, status, fix in checks:
        cell(ws, f"B{r}", label, F(9), None, LEFT, None, BORDER)
        cell(ws, f"C{r}", val, F(9), None, RIGHT, "#,##0", BORDER)
        cell(ws, f"D{r}", status, F(9, True), None, CENTER, None, BORDER)
        ws.merge_cells(f"E{r}:L{r}")
        cell(ws, f"E{r}", fix, F(9), None, LEFT, None, BORDER)
        r += 1
    last = r - 1
    cell(ws, f"B{r + 1}", "종합 상태", F(12, True))
    cell(ws, f"C{r + 1}", f'=IF(COUNTIF(D{first}:D{last},"오류")>0,"오류",IF(COUNTIF(D{first}:D{last},"확인")>0,"확인 필요","정상"))',
         F(12, True), None, CENTER, None, BORDER)
    name(wb, "CHECK_STATUS", f"실행!$C${r + 1}")
    for rng in (f"D{first}:D{last}", f"C{r + 1}"):
        for val, bg, fg in (("정상", "C6EFCE", "006100"), ("오류", "FFC7CE", "9C0006"),
                            ("확인", "FFEB9C", "9C5700"), ("확인 필요", "FFEB9C", "9C5700")):
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                          fill=PatternFill("solid", fgColor=bg), font=Font(color=fg, bold=True)))
    cell(ws, f"B{r + 3}", "※ 매크로 최초 1회 설치: Alt+F11 → 파일 → 파일 가져오기 → report_macro.bas → 닫기 → "
                          "'Excel 매크로 사용 통합 문서(*.xlsm)'로 저장. 이후 Alt+F8로 실행.", F(9, color="595959"))


# ---------------------------------------------------------------------------
# 담보별집계 (담보코드 단위)
# ---------------------------------------------------------------------------
def sheet_code(wb):
    ws = wb.create_sheet(CD)
    heads = ["담보코드", "급부(담보명)", "대급부명", "구분", "고당대통",
             "UY당기 보험료", "UY당기 손해액", "UY비교 보험료", "UY비교 손해액",
             "누적 보험료", "누적 손해액", "전월누적 보험료", "전월누적 손해액", "누적 건수",
             "급부 번호", "로데이터 행수", "중복"]
    for j, h in enumerate(heads, start=1):
        cell(ws, f"{get_column_letter(j)}1", h, F(9, True), FILL_HDR, CENTER, None, BORDER)
    widths(ws, [10, 30, 11, 9, 7] + [13] * 9 + [8, 9, 6])
    ws.row_dimensions[1].height = 30
    # code_master 는 열 전체 + ROW() 로 참조 → 마스터에서 행을 삭제·삽입해도 #REF! 가 나지 않음
    mc = lambda col: f"INDEX({MST}!${col}:${col},ROW())"
    for r in range(2, CL + 1):
        ea, eb = esc(f"$A{r}"), esc(f"$B{r}")
        g = f"{R}$G:$G,{ea}"
        ws[f"A{r}"] = f'=IF(TRIM({mc("A")}&"")="","",{mc("A")})'
        ws[f"B{r}"] = f'=IF($A{r}="","",TRIM({mc("B")}&""))'
        ws[f"C{r}"] = f'=IF($A{r}="","",TRIM({mc("C")}&""))'
        ws[f"D{r}"] = f'=IF($A{r}="","",IF(TRIM({mc("D")}&"")="",C{r},TRIM({mc("D")}&"")))'
        ws[f"E{r}"] = f'=IF($A{r}="","",IF(UPPER(TRIM({mc("E")}&""))="Y",1,0))'
        for col, w in zip("FGHIJKLM", ["UC", "UC", "UCC", "UCC", "ALL", "ALL", "ALLP", "ALLP"]):
            expr = prem(w, "," + g) if col in "FHJL" else loss(w, "," + g)
            ws[f"{col}{r}"] = f'=IF($A{r}="","",{expr})'
            ws[f"{col}{r}"].number_format = "#,##0"
        ws[f"N{r}"] = f'=IF($A{r}="","",SUMIFS({R}$K:$K,{CRIT["ALL"]},{g}))'
        ws[f"O{r}"] = f'=IF(OR($A{r}="",$B{r}=""),"",IF(COUNTIF($B$2:$B{r},{eb})=1,MAX($O$1:$O{r - 1})+1,""))'
        ws[f"P{r}"] = f'=IF($A{r}="","",COUNTIF({R}$G:$G,{ea}))'
        ws[f"Q{r}"] = f'=IF($A{r}="","",IF(COUNTIF($A$2:$A${CL},{ea})>1,1,0))'
    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# 급부별집계 (담보명 단위) + 표 선정
# ---------------------------------------------------------------------------
GB_HEADS = {
    "A": "급부", "B": "구분(표1·2)", "C": "대급부(표3)", "D": "고당대통",
    "E": "UY당기 보험료", "F": "UY당기 손해액", "G": "UY비교 보험료", "H": "UY비교 손해액",
    "I": "누적 보험료", "J": "누적 손해액", "K": "전월누적 보험료", "L": "전월누적 손해액", "M": "누적 건수",
    "N": "UY당기 손해율", "O": "UY비교 손해율", "P": "증감", "Q": "누적 손해율", "R": "전월누적 손해율",
    "S": "누적 P/F", "T": "순위대상", "U": "손해율키", "V": "손해율순위", "W": "손해율키(고당대통제외)",
    "X": "손해율순위(고당대통제외)", "Y": "표1 포함", "Z": "표1 구분", "AA": "표1 정렬키", "AB": "표1 순번",
    "AC": "표1 색(1신규만악화 2누적만악화)", "AD": "표2 포함", "AE": "표2 정렬키", "AF": "표2 순번",
    "AG": "표2 색(1악화 2감소)", "AH": "보험료키", "AI": "보험료순위", "AJ": "표3 정렬키", "AK": "표3 순번",
}


def sheet_gb(wb):
    ws = wb.create_sheet(GB)
    for col, h in GB_HEADS.items():
        cell(ws, f"{col}1", h, F(9, True), FILL_HDR, CENTER, None, BORDER)
    widths(ws, [32, 9, 11, 7] + [13] * 9 + [9] * 6 + [8] * 18)
    ws.row_dimensions[1].height = 42
    rng = lambda col: f"${col}$2:${col}${GL}"
    cb = lambda col: f"{C_}${col}$2:${col}${CL}"
    asc_rank = lambda key, r: f'=IF({key}{r}="","",RANK({key}{r},{rng(key)},1)+COUNTIF(${key}$2:{key}{r},{key}{r})-1)'
    desc_rank = lambda key, r: f'=IF({key}{r}="","",RANK({key}{r},{rng(key)},0)+COUNTIF(${key}$2:{key}{r},{key}{r})-1)'
    for r in range(2, GL + 1):
        ea = esc(f"$A{r}")
        f = {
            "A": f'=IFERROR(INDEX({cb("B")},MATCH(ROW()-1,{cb("O")},0)),"")',
            "B": f'=IF($A{r}="","",INDEX({cb("D")},MATCH({ea},{cb("B")},0)))',
            "C": f'=IF($A{r}="","",INDEX({cb("C")},MATCH({ea},{cb("B")},0)))',
            "D": f'=IF($A{r}="","",IF(COUNTIFS({cb("B")},{ea},{cb("E")},1)>0,1,0))',
        }
        for dst, src in zip("EFGHIJKLM", "FGHIJKLMN"):
            f[dst] = f'=IF($A{r}="","",SUMIF({cb("B")},{ea},{cb(src)}))'
        f.update({
            "N": f'=IF(N(E{r})=0,"",F{r}/E{r})',
            "O": f'=IF(N(G{r})=0,"",H{r}/G{r})',
            "P": f'=IF(OR(N{r}="",O{r}=""),"",N{r}-O{r})',
            "Q": f'=IF(N(I{r})=0,"",J{r}/I{r})',
            "R": f'=IF(N(K{r})=0,"",L{r}/K{r})',
            "S": f'=IF(N(I{r})=0,"",I{r}/SUM({rng("I")}))',
            "T": f'=IF(AND(N(E{r})>0,N(E{r})>=MIN_SHARE*SUM({rng("E")})),1,0)',
            "U": f'=IF(T{r}=1,N{r},"")',
            "V": desc_rank("U", r),
            "W": f'=IF(AND(T{r}=1,N(D{r})=0),N{r},"")',
            "X": desc_rank("W", r),
            "Y": f'=IF($A{r}="",0,IF(OR(AND(X{r}<>"",N(X{r})<=TOP_LR_N),N(D{r})=1),1,0))',
            "Z": f'=IF($A{r}="","",IF(N(D{r})=1,GODANG_GROUP,B{r}))',
            "AA": f'=IF(Y{r}=1,IFERROR(MATCH({esc(f"Z{r}")},ORDER1,0),50)*100000+IF(V{r}="",99999,V{r}),"")',
            "AB": asc_rank("AA", r),
            # 표1 색: 1 빨강 = UY 신규 악화 & 누적 악화 아님 / 2 파랑 = UY 신규 양호 & 누적만 악화 / 둘 다 악화·둘 다 양호 = 0
            "AC": (f'=IF(Y{r}<>1,"",IF(AND(N(P{r})>0,NOT(AND(Q{r}<>"",R{r}<>"",N(Q{r})>N(R{r})))),1,'
                   f'IF(AND(N(P{r})<=0,AND(Q{r}<>"",R{r}<>"",N(Q{r})>N(R{r}))),2,0)))'),
            "AD": f'=IF(AND($A{r}<>"",Y{r}=0,N{r}<>""),IF(N{r}>OVER_LR,1,0),0)',
            "AE": f'=IF(AD{r}=1,IFERROR(MATCH({esc(f"B{r}")},ORDER1,0),50)*100000+IF(V{r}="",99999,V{r}),"")',
            "AF": asc_rank("AE", r),
            "AG": f'=IF(AD{r}<>1,"",IF(N(P{r})>0,1,IF(N(P{r})<0,2,0)))',
            "AH": f'=IF(N(E{r})>0,E{r},"")',
            "AI": desc_rank("AH", r),
            "AJ": f'=IF(AND(AI{r}<>"",N(AI{r})<=TOP_PREM_N),IFERROR(MATCH({esc(f"C{r}")},ORDER3,0),50)*100000+AI{r},"")',
            "AK": asc_rank("AJ", r),
        })
        for col, v in f.items():
            ws[f"{col}{r}"] = v
        for col in "EFGHIJKL":
            ws[f"{col}{r}"].number_format = "#,##0"
        for col in "NOPQRS":
            ws[f"{col}{r}"].number_format = "0.0%"
    ws.freeze_panes = "B2"


# ---------------------------------------------------------------------------
# 보고서 시트 (A열 태그 → VBA가 위에서부터 읽어 워드 생성)
# ---------------------------------------------------------------------------
TAG_FONT = F(8, color="A6A6A6")
STYLE = {  # 태그: 엑셀 화면 글꼴
    "H1": F(15, True), "H2": F(13, True), "BUL": F(11), "CAP": F(11, True),
    "UNIT": F(9), "LEG": F(9, color="7F7F7F"), "NOTE": F(9),
}


class Report:
    def __init__(self, wb):
        self.wb = wb
        self.ws = wb.create_sheet(RPT, 1)
        ws = self.ws
        widths(ws, [16, 12, 36, 13, 13, 13, 13, 13, 13, 5, 5])
        ws.column_dimensions["J"].hidden = True      # 색 코드
        ws.column_dimensions["K"].hidden = True      # 급부별집계 행 위치
        ws.sheet_view.showGridLines = False
        cell(ws, "A1", "← 워드 출력 태그 (수정 금지)", TAG_FONT)
        self.r = 3

    def text(self, tag, value, gap_before=0):
        self.r += gap_before
        cell(self.ws, f"A{self.r}", tag, TAG_FONT)
        cell(self.ws, f"B{self.r}", value, STYLE[tag])
        self.r += 1

    def pagebreak(self):
        self.r += 1
        cell(self.ws, f"A{self.r}", "PB", TAG_FONT)
        cell(self.ws, f"B{self.r}", "─" * 20 + " (워드 페이지 나눔) " + "─" * 20, F(8, color="BFBFBF"))
        self.r += 2

    def tag_table(self, head, body, foot, cols, group, color, word_widths):
        cell(self.ws, f"A{self.r}", f"TBL,{head},{body},{foot},{cols},{group},{color},"
                                    + ";".join(str(w) for w in word_widths), TAG_FONT)


def table_static_cy(rp: Report):
    ws, r = rp.ws, rp.r
    rp.tag_table(4, 0, 0, 6, 0, 0, [7, 9, 13, 13, 12, 17])
    merge(ws, f"B{r}:C{r}", "구 분")
    ws[f"D{r}"] = '="CY "&L_BASE&"월"'
    ws[f"E{r}"] = '="CY "&L_PREV&"월"'
    ws[f"F{r}"] = "증감"
    ws[f"G{r}"] = "누적"
    box(ws, f"B{r}:G{r}", FILL_HDR, F(10, True))
    merge(ws, f"B{r + 1}:B{r + 3}", "손해율")
    for i, t in enumerate(["", "위험 P", "손해액"]):
        ws[f"C{r + 1 + i}"] = t
    p, l = r + 2, r + 3
    for col, w in (("D", "CY"), ("E", "CYP"), ("G", "ALL")):
        ws[f"{col}{p}"] = f"={prem(w)}"
        ws[f"{col}{l}"] = f"={loss(w)}"
    for col in "DE":
        ws[f"{col}{r + 1}"] = f'=IFERROR({col}{l}/{col}{p},"")'
    ws[f"G{r + 1}"] = (f'=IFERROR(TEXT(G{l}/G{p},"0%"),"-")&CHAR(10)&"(*~"&MID(L_PREV,2,9)&"월: "&'
                       f'IFERROR(TEXT({loss("ALLP")}/{prem("ALLP")},"0%"),"-")&")"')
    ws[f"F{r + 1}"] = f'=IF(OR(D{r + 1}="",E{r + 1}=""),"",D{r + 1}-E{r + 1})'
    ws[f"F{p}"] = f"=D{p}-E{p}"
    ws[f"F{l}"] = f"=D{l}-E{l}"
    box(ws, f"B{r + 1}:G{r + 3}", font=F(10))
    ws[f"B{r + 1}"].font = F(10, True)
    for col in "DE":
        ws[f"{col}{r + 1}"].number_format = NF_LR
    ws[f"F{r + 1}"].number_format = NF_PT
    for rr in (p, l):
        for col in "DEFG":
            ws[f"{col}{rr}"].number_format = NF_AMT
    ws.row_dimensions[r + 1].height = 30
    name(rp.wb, "CY_LR", f"{RPT}!$D${r + 1}")
    name(rp.wb, "CY_LR_DIFF", f"{RPT}!$F${r + 1}")
    rp.r = r + 4


def table_static_uy(rp: Report):
    ws, r = rp.ws, rp.r
    rp.tag_table(5, 0, 0, 8, 0, 0, [7, 8, 10, 10, 9, 11, 13, 9])
    merge(ws, f"B{r}:C{r + 1}", "구 분")
    merge(ws, f"D{r}:F{r}", '="UY 직전 "&N_MON&"개월 당월손해율"')
    merge(ws, f"G{r}:I{r}", '="UY 직전 "&N_MON&"개월 누적손해율"')
    for col, v in zip("DEFGHI", ["=L_BASE", "=L_PREV", "比", "=L_UYC", "=L_CMP", "比"]):
        ws[f"{col}{r + 1}"] = v
    box(ws, f"B{r}:I{r + 1}", FILL_HDR, F(10, True))
    b = r + 2
    merge(ws, f"B{b}:B{b + 2}", "손해율")
    for i, t in enumerate(["", "위험 P", "손해액"]):
        ws[f"C{b + i}"] = t
    p, l = b + 1, b + 2
    for col, w in (("D", "UM"), ("E", "UMC"), ("G", "UC"), ("H", "UCC")):
        ws[f"{col}{p}"] = f"={prem(w)}"
        ws[f"{col}{l}"] = f"={loss(w)}"
        ws[f"{col}{b}"] = f'=IFERROR({col}{l}/{col}{p},"")'
    for d, a, c in (("F", "D", "E"), ("I", "G", "H")):
        ws[f"{d}{b}"] = f'=IF(OR({a}{b}="",{c}{b}=""),"",{a}{b}-{c}{b})'
        ws[f"{d}{p}"] = f"={a}{p}-{c}{p}"
        ws[f"{d}{l}"] = f"={a}{l}-{c}{l}"
    box(ws, f"B{b}:I{b + 2}", font=F(10))
    ws[f"B{b}"].font = F(10, True)
    for col in "DEGH":
        ws[f"{col}{b}"].number_format = NF_LR
    for col in "FI":
        ws[f"{col}{b}"].number_format = NF_PT
    for rr in (p, l):
        for col in "DEFGHI":
            ws[f"{col}{rr}"].number_format = NF_AMT
    for nm, ref in (("UM_LR", f"D{b}"), ("UM_DIFF", f"F{b}"), ("UC_LR", f"G{b}"), ("UCC_LR", f"H{b}"), ("UC_DIFF", f"I{b}")):
        name(rp.wb, nm, f"{RPT}!${ref[0]}${ref[1:]}")
    rp.r = b + 3


def gb_header(ws, r):
    """표(1)(2) 공통 2행 헤더."""
    merge(ws, f"B{r}:B{r + 1}", "구분")
    merge(ws, f"C{r}:C{r + 1}", "급부")
    merge(ws, f"D{r}:F{r}", '="손해율_"&N_MON&"개월 동기간 비교"')
    merge(ws, f"G{r}:G{r + 1}", "누적\n손해율")
    ws[f"D{r + 1}"] = f'="UY"&YYN&CHAR(10)&{mm("UY_START")}&"~"&{mm("BASE_YM")}&"월"'
    ws[f"E{r + 1}"] = f'="UY"&{yy("CMP_START")}&"."&{mm("CMP_START")}&CHAR(10)&"~"&{yy("CMP_END")}&"."&{mm("CMP_END")}&"월"'
    ws[f"F{r + 1}"] = "증감"
    box(ws, f"B{r}:G{r + 1}", FILL_HDR, F(10, True))
    ws.row_dimensions[r + 1].height = 30


def gb_body(rp: Report, r0, n, slot_col, group_col, color_col, cols):
    """급부별 표 본문 슬롯 n행. cols: (보고서 열, 급부별집계 열, 서식, 정렬)."""
    ws = rp.ws
    g = lambda col: f"{G_}${col}$2:${col}${GL}"
    for k in range(1, n + 1):
        r = r0 + k - 1
        ws[f"K{r}"] = f'=IFERROR(MATCH({k},{g(slot_col)},0),"")'
        if color_col:
            ws[f"J{r}"] = f'=IF($K{r}="","",INDEX({g(color_col)},$K{r}))'
        c = ws[f"B{r}"]
        c.value = f'=IF($K{r}="","",INDEX({g(group_col)},$K{r}))'
        c.alignment, c.border, c.font = CENTER, BORDER, F(10)
        for col, src, fmt, al in cols:
            c = ws[f"{col}{r}"]
            c.value = f'=IF($K{r}="","",INDEX({g(src)},$K{r}))'
            c.number_format, c.alignment, c.border, c.font = fmt or "General", al, BORDER, F(10)
    if color_col:
        rng = f"B{r0}:G{r0 + n - 1}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"$J{r0}=1"], font=Font(color="C00000")))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"$J{r0}=2"], font=Font(color="0070C0")))


GB_COLS_12 = [("C", "A", None, LEFT), ("D", "N", NF_LR, CENTER), ("E", "O", NF_LR, CENTER),
              ("F", "P", NF_PT, CENTER), ("G", "Q", NF_LR, CENTER)]


def table_gb1(rp: Report):
    ws, r = rp.ws, rp.r
    n = SLOTS["T1"]
    rp.tag_table(2, n, 3, 6, 1, 9, [7, 24, 9, 10, 8, 9])
    gb_header(ws, r)
    b = r + 2
    gb_body(rp, b, n, "AB", "Z", "AC", GB_COLS_12)
    f = b + n
    Y = f"{G_}$Y$2:$Y${GL}"
    s = lambda col: f"SUMIF({Y},1,{G_}${col}$2:${col}${GL})"
    foot = [
        ("주요급부 평균(합계)", f'=IFERROR({s("F")}/{s("E")},"")', f'=IFERROR({s("H")}/{s("G")},"")',
         f'=IFERROR({s("J")}/{s("I")},"")', NF_LR, NF_PT),
        ("전급부 평균(합계)", f"={lr('UC')}", f"={lr('UCC')}", f"={lr('ALL')}", NF_LR, NF_PT),
        ("전급부 대비 위험보험료 비중", f'=IFERROR({s("E")}/{prem("UC")},"")', f'=IFERROR({s("G")}/{prem("UCC")},"")',
         f'=IFERROR({s("I")}/{prem("ALL")},"")', NF_SH, NF_SH_PT),
    ]
    for i, (label, d, e, gg, fmt, dfmt) in enumerate(foot):
        rr = f + i
        merge(ws, f"B{rr}:C{rr}", label)
        ws[f"D{rr}"], ws[f"E{rr}"], ws[f"G{rr}"] = d, e, gg
        ws[f"F{rr}"] = f'=IF(OR(D{rr}="",E{rr}=""),"",D{rr}-E{rr})'
        box(ws, f"B{rr}:G{rr}", font=F(10, True))
        for col in "DEG":
            ws[f"{col}{rr}"].number_format = fmt
        ws[f"F{rr}"].number_format = dfmt
    rp.r = f + 3


def table_gb2(rp: Report):
    ws, r = rp.ws, rp.r
    n = SLOTS["T2"]
    rp.tag_table(2, n, 0, 6, 1, 9, [7, 24, 9, 10, 8, 9])
    gb_header(ws, r)
    gb_body(rp, r + 2, n, "AF", "B", "AG", GB_COLS_12)
    rp.r = r + 2 + n


def table_gb3(rp: Report):
    ws, r = rp.ws, rp.r
    n = SLOTS["T3"]
    rp.tag_table(1, n, 0, 6, 1, 0, [8, 22, 10, 10, 7, 9])
    for col, h in zip("BCDEFG", ["급부", "세부급부", "누적\n위험보험료", "누적\n위험보험료 P/F", "건수", "누적\n손해율"]):
        ws[f"{col}{r}"] = h
    box(ws, f"B{r}:G{r}", FILL_HDR, F(10, True))
    ws.row_dimensions[r].height = 30
    cols = [("C", "A", None, LEFT), ("D", "I", NF_AMT, CENTER), ("E", "S", NF_SH, CENTER),
            ("F", "M", NF_CNT, CENTER), ("G", "Q", NF_LR, CENTER)]
    gb_body(rp, r + 1, n, "AK", "C", None, cols)
    rp.r = r + 1 + n


def sheet_report(wb):
    rp = Report(wb)
    N = "N_MON"
    rp.text("H1", "2. 손해율 현황")
    rp.text("H2", "2.1 건강보험 손해율 현황")
    rp.text("BUL", '=IFERROR("○ CY"&L_BASE&"월 건강보험 손해율은 "&TEXT(CY_LR,"0%")&"로 CY"&L_PREV&"월 대비 "'
                   '&TEXT(ABS(CY_LR_DIFF),"0%")&"p "&IF(CY_LR_DIFF<0,"감소함","증가함"),"○ (데이터 없음)")')
    rp.text("CAP", '="[CY 기준 건강보험 손해율 _ "&L_BASE&"월 마감]"', 1)
    rp.text("UNIT", "(단위 : %, 백만원)")
    table_static_cy(rp)
    rp.text("NOTE", '="※ CY"&L_BASE&"월 : 계약체결월도와 관계없이 "&L_BASE&"월에 유입된 위험P와 지급된 보험금을 바탕으로 산출한 손해율"')
    rp.text("NOTE", '="    누적 : 최초판매시점("&L_LAUNCH&"월)부터 "&L_BASE&"월까지 유입된 위험P와 지급된 보험금을 바탕으로 산출한 손해율"')
    rp.text("NOTE", "    손해액 : 보험금 및 비금변동")

    rp.text("BUL", '=IFERROR("○ UY"&L_BASE&"월 건강보험 손해율은 "&TEXT(UM_LR,"0%")&"로 동 기간 대비 "&TEXT(ABS(UM_DIFF),"0%")&"p "'
                   '&IF(UM_DIFF<=0,"개선","악화")&"되었으며, 동 기간 누적 손해율로는 "&TEXT(ABS(UC_DIFF),"0%")&"p ("'
                   '&TEXT(UCC_LR,"0%")&" → "&TEXT(UC_LR,"0%")&") "&IF(UC_DIFF<=0,"개선","악화")&"됨","○ (데이터 없음)")', 1)
    rp.text("CAP", '="[UY\'"&YYN&" 기준 건강보험 손해율 _ "&L_BASE&"월 마감]"', 1)
    rp.text("UNIT", "(단위 : %, 백만원)")
    table_static_uy(rp)
    rp.text("NOTE", f'="※ UY 직전 "&{N}&"개월 당월손해율 "&L_BASE&" : "&L_UYC&"월에 가입된 계약의 "&L_BASE&"월 마감 손해율"')
    rp.text("NOTE", f'="    UY 직전 "&{N}&"개월 당월손해율 "&L_PREV&" : "&L_CMP&"월에 가입된 계약의 "&L_PREV&"월 마감 손해율"')
    rp.text("NOTE", f'="    UY 직전 "&{N}&"개월 누적손해율 "&L_UYC&" : "&L_UYC&"월에 가입된 계약의 "&L_BASE&"월까지 누계 손해율"')
    rp.text("NOTE", f'="    UY 직전 "&{N}&"개월 누적손해율 "&L_CMP&" : "&L_CMP&"월에 가입된 계약의 "&L_PREV&"월까지 누계 손해율"')

    rp.pagebreak()
    rp.text("H1", '="별첨 1. "&L_BASE&"월 급부별 손해율"')
    rp.text("CAP", '="[급부별 건강보험 손해율 현황 _ (1) 주요급부별 손해율(UY"&YYN&" 손해율 상위 "&TOP_LR_N&" + 고당대통)]"', 1)
    rp.text("LEG", '="* UY"&YYN&" 신규로 악화된 담보 / 누적손해율만 악화된 담보"')
    rp.text("UNIT", "(단위 : %, %p)")
    table_gb1(rp)
    uyc = f'"UY"&YYN&"."&{mm("UY_START")}&"~"&{mm("BASE_YM")}&"월"'
    cmp_ = f'"UY"&{yy("CMP_START")}&"."&{mm("CMP_START")}&"~"&{yy("CMP_END")}&"."&{mm("CMP_END")}&"월"'
    rp.text("NOTE", f'="※ "&{uyc}&" 손해율: "&L_UYC&"월에 가입된 계약의 "&L_UYC&"월 누계 손해율"')
    rp.text("NOTE", f'="    "&{cmp_}&" 손해율: "&L_CMP&"월에 가입된 계약의 "&L_CMP&"월 누계 손해율"')

    rp.pagebreak()
    rp.text("CAP", '="[급부별 건강보험 손해율 현황 _ (2) 그외 "&TEXT(OVER_LR,"0%")&" 초과 급부(UY"&YYN&" 손해율 "'
                   '&TEXT(OVER_LR,"0%")&"초과 대상)]"')
    rp.text("LEG", f'="* 직전 "&{N}&"개월대비 손해율 악화 담보 / 직전 "&{N}&"개월대비 손해율 감소 담보"')
    rp.text("UNIT", "(단위 : %, %p)")
    table_gb2(rp)

    rp.pagebreak()
    rp.text("CAP", '="[급부별 건강보험 손해율 현황 _ (3) 위험보험료 규모 상위 "&TOP_PREM_N&" 급부(UY"&YYN&" 기준 상위 "'
                   '&TOP_PREM_N&" 대상)]"')
    rp.text("UNIT", "(단위 : 백만원, 건, %)")
    table_gb3(rp)
    rp.r += 1
    cell(rp.ws, f"A{rp.r}", "END", TAG_FONT)


# ---------------------------------------------------------------------------
RAW_COLS = ["계약년월", "마감년월", "주상품코드", "특약상품코드", "경과개월수", "심사유형(일반.간편)",
            "담보코드", "위험보험료 합계", "지급보험금합계", "지급비금변동분 합계", "건수"]
MST_COLS = ["담보코드", "담보명", "대급부명", "구분", "고당대통"]


def sheet_raw(wb, raw: pd.DataFrame):
    ws = wb.create_sheet(RAW)
    ws.append(RAW_COLS)
    for rec in raw.itertuples(index=False):
        ws.append(list(rec))
    for j in range(1, len(RAW_COLS) + 1):
        c = ws.cell(1, j)
        c.font, c.fill, c.alignment, c.border = F(9, True), FILL_HDR, CENTER, BORDER
    widths(ws, [9, 9, 10, 11, 8, 12, 10, 14, 13, 14, 7])
    ws.freeze_panes = "A2"


def sheet_master(wb, master: pd.DataFrame):
    ws = wb.create_sheet(MST)
    ws.append(MST_COLS)
    for rec in master.itertuples(index=False):
        ws.append([None if v == "" else v for v in rec])
    for j in range(1, len(MST_COLS) + 1):
        c = ws.cell(1, j)
        c.font, c.fill, c.alignment, c.border = F(9, True), FILL_HDR, CENTER, BORDER
    widths(ws, [10, 36, 12, 10, 9, 3, 60])
    ws.freeze_panes = "A2"
    cell(ws, "G1", "D열 구분: 표(1)(2) 구분 (빈칸이면 대급부명) / E열 고당대통: Y 표시 / 급부 = 담보명 기준 합산",
         F(9, color="595959"))


def sheet_macro(wb):
    """VBA 코드 사본 (.bas/.txt 반입이 막힐 때 복사·붙여넣기용)."""
    ws = wb.create_sheet("매크로코드")
    widths(ws, [120])
    cell(ws, "A1", "매크로 설치(파일 반입이 막힐 때): 아래 A3부터 마지막 줄까지 선택 → 복사 → Alt+F11 → 삽입 → 모듈 → 붙여넣기 → "
                   ".xlsm으로 저장", F(10, True, "C00000"))
    lines = (BASE / "excel" / "report_macro.bas").read_text(encoding="cp949").splitlines()
    for i, line in enumerate([l for l in lines if not l.startswith("Attribute ")], start=3):
        c = ws.cell(i, 1)
        c.value = line if line else None
        c.data_type = "s"
        c.font = Font(name="Consolas", size=10)


def main():
    raw = pd.read_csv(BASE / "data" / "raw_data.csv",
                      dtype={"주상품코드": str, "특약상품코드": str, "담보코드": str})
    master = pd.read_csv(BASE / "data" / "code_master.csv", dtype=str, keep_default_na=False)
    wb = Workbook()
    sheet_run(wb)
    sheet_report(wb)
    sheet_gb(wb)
    sheet_code(wb)
    sheet_raw(wb, raw[RAW_COLS])
    sheet_master(wb, master[MST_COLS])
    sheet_macro(wb)
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"[OK] {OUT}")


if __name__ == "__main__":
    main()
